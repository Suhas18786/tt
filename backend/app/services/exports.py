"""PDF and Excel export of timetable grids."""
from __future__ import annotations

import io
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..scheduler.context import DAY_NAMES

VIEW_KEYS = {"batch": ("batch_id", "batch_code"), "faculty": ("faculty_id", "faculty_name"),
             "room": ("room_id", "room_label"), "department": ("department_id", "department_code")}


def group_sessions(sessions: list, view: str, entity_id: int | None = None) -> list:
    """-> [(group_id, title, [sessions])] sorted by title."""
    if view not in VIEW_KEYS:
        raise ValueError(f"unknown view '{view}'")
    key, title_key = VIEW_KEYS[view]
    groups, titles = defaultdict(list), {}
    for s in sessions:
        if entity_id is not None and s[key] != entity_id:
            continue
        groups[s[key]].append(s)
        titles[s[key]] = s[title_key]
    return sorted(((g, titles[g], v) for g, v in groups.items()), key=lambda x: x[1])


def cell_text(s: dict, view: str) -> str:
    if view == "batch":
        return f"{s['subject_code']}\n{s['faculty_name']}\n{s['room_label']}"
    if view == "faculty":
        return f"{s['subject_code']}\n{s['batch_code']}\n{s['room_label']}"
    if view == "room":
        return f"{s['subject_code']}\n{s['batch_code']}\n{s['faculty_name']}"
    return f"{s['subject_code']}\n{s['batch_code']}\n{s['faculty_name']}\n{s['room_label']}"


def build_grid(config: dict, sessions: list, view: str):
    """Returns (header, rows, spans). rows[i] = [period label, cell per working day]."""
    days = config["working_days"]
    periods = config["periods"]
    start_cells = {}
    covered = set()
    for s in sessions:
        start_cells[(s["day"], s["start_period"])] = s
        for q in range(s["start_period"] + 1, s["start_period"] + s["duration"]):
            covered.add((s["day"], q))
    header = ["Period"] + [DAY_NAMES[d] for d in days]
    rows, spans = [], []
    for r, p in enumerate(periods):
        row = [f"{p['start']}-{p['end']}" + ("" if p.get("kind", "class") == "class" else f"\n({p['kind']})")]
        for c, d in enumerate(days):
            if p.get("kind", "class") != "class" or (d, p["index"]) in set(map(tuple, config.get("blocked_slots", []))):
                row.append("")
            elif (d, p["index"]) in start_cells:
                s = start_cells[(d, p["index"])]
                row.append(cell_text(s, view))
                if s["duration"] > 1:
                    spans.append((c + 1, r, c + 1, r + s["duration"] - 1))
            else:
                row.append("" if (d, p["index"]) in covered else "")
        rows.append(row)
    return header, rows, spans


def _sheet_name(title: str, used: set) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "_", title)[:31] or "Sheet"
    base, i = name, 2
    while name in used:
        name = f"{base[:28]}_{i}"
        i += 1
    used.add(name)
    return name


def render_xlsx(gen: dict, view: str = "batch", entity_id: int | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    thin = Side(style="thin", color="999999")
    groups = group_sessions(gen["sessions"], view, entity_id)
    if not groups:
        ws = wb.create_sheet("Timetable")
        ws["A1"] = "No scheduled sessions"
    for _gid, title, sess in groups:
        ws = wb.create_sheet(_sheet_name(title, used))
        header, rows, spans = build_grid(gen["config"], sess, view)
        ws["A1"] = f"{gen['name']} - {view.title()}: {title}"
        ws["A1"].font = Font(bold=True, size=13)
        for c, h in enumerate(header, start=1):
            cell = ws.cell(row=3, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r, row in enumerate(rows):
            for c, val in enumerate(row, start=1):
                cell = ws.cell(row=4 + r, column=c, value=val)
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if val and c > 1:
                    cell.fill = PatternFill("solid", fgColor="DCE9F7")
            ws.row_dimensions[4 + r].height = 52
            if gen["config"]["periods"][r].get("kind", "class") != "class":
                for c in range(1, len(header) + 1):
                    ws.cell(row=4 + r, column=c).fill = PatternFill("solid", fgColor="EEEEEE")
        for c1, r1, c2, r2 in spans:
            ws.merge_cells(start_row=4 + r1, start_column=c1 + 1, end_row=4 + r2, end_column=c2 + 1)
        ws.column_dimensions["A"].width = 16
        for c in range(2, len(header) + 1):
            ws.column_dimensions[ws.cell(row=3, column=c).column_letter].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_pdf(gen: dict, view: str = "batch", entity_id: int | None = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=24, bottomMargin=20,
                            title=gen["name"])
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell")
    cell_style.fontSize = 7
    cell_style.leading = 8.5
    story = []
    groups = group_sessions(gen["sessions"], view, entity_id)
    if not groups:
        story.append(Paragraph("No scheduled sessions", styles["Heading2"]))
    for i, (_gid, title, sess) in enumerate(groups):
        header, rows, spans = build_grid(gen["config"], sess, view)
        story.append(Paragraph(f"{gen['name']} &mdash; {view.title()}: {title}", styles["Heading2"]))
        story.append(Spacer(1, 6))
        data = [header] + [[Paragraph(c.replace("\n", "<br/>"), cell_style) if c else "" for c in row] for row in rows]
        t = Table(data, repeatRows=1, colWidths=[70] + [(landscape(A4)[0] - 40 - 70) / (len(header) - 1)] * (len(header) - 1))
        style = [("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                 ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 8)]
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                if c > 0 and val:
                    style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor("#DCE9F7")))
            if gen["config"]["periods"][r - 1].get("kind", "class") != "class":
                style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#EEEEEE")))
        for c1, r1, c2, r2 in spans:
            style.append(("SPAN", (c1, r1 + 1), (c2, r2 + 1)))
        t.setStyle(TableStyle(style))
        story.append(t)
        if i < len(groups) - 1:
            story.append(PageBreak())
    doc.build(story)
    return buf.getvalue()
