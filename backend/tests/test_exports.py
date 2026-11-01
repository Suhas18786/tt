"""Excel / PDF export (needs openpyxl + reportlab only)."""
import io

from openpyxl import load_workbook

from app.scheduler.context import DEFAULT_PERIODS
from app.services.exports import build_grid, group_sessions, render_pdf, render_xlsx


def sample_gen():
    base = dict(subject_name="x", room_type="lecture_room", department_id=1,
                department_code="CSE", semester_id=1, assignment_id=1)
    sessions = [
        dict(base, id=1, batch_id=1, batch_code="B1", subject_id=1, subject_code="CS101", faculty_id=1,
             faculty_name="Dr. A", room_id=1, room_label="LH-1", day=0, start_period=0, duration=1, subject_type="lecture"),
        dict(base, id=2, batch_id=1, batch_code="B1", subject_id=2, subject_code="CS102L", faculty_id=2,
             faculty_name="Dr. B", room_id=2, room_label="CL-1", day=1, start_period=1, duration=2, subject_type="lab"),
        dict(base, id=3, batch_id=2, batch_code="B2", subject_id=1, subject_code="CS101", faculty_id=1,
             faculty_name="Dr. A", room_id=1, room_label="LH-1", day=2, start_period=5, duration=1, subject_type="lecture"),
    ]
    return {"name": "Test TT", "config": {"working_days": [0, 1, 2, 3, 4], "periods": DEFAULT_PERIODS, "blocked_slots": []},
            "sessions": sessions}


def test_grid_marks_spans_for_multi_period_sessions():
    gen = sample_gen()
    _h, rows, spans = build_grid(gen["config"], [s for s in gen["sessions"] if s["batch_id"] == 1], "batch")
    assert (2, 1, 2, 2) in spans              # day column 2, periods 2-3 merged
    assert "CS101" in rows[0][1] and "Dr. A" in rows[0][1] and "LH-1" in rows[0][1]


def test_group_sessions_by_each_view():
    gen = sample_gen()
    assert [t for _i, t, _s in group_sessions(gen["sessions"], "batch")] == ["B1", "B2"]
    assert [t for _i, t, _s in group_sessions(gen["sessions"], "faculty")] == ["Dr. A", "Dr. B"]
    assert len(group_sessions(gen["sessions"], "room", entity_id=1)) == 1


def test_xlsx_export_has_one_sheet_per_batch_and_merged_lab_block():
    wb = load_workbook(io.BytesIO(render_xlsx(sample_gen(), "batch")))
    assert wb.sheetnames == ["B1", "B2"]
    ws = wb["B1"]
    assert ws["B4"].value.startswith("CS101")
    # Mon Class=B, Mon Lab=C, Tue Class=D, Tue Lab=E: the 2-period lab block sits merged in Tue Lab
    assert ws["E3"].value == "Tuesday Lab" or "Lab" in str(ws["E3"].value)
    assert ws["E5"].value.startswith("CS102L")
    assert any("E5:E6" in str(r) for r in ws.merged_cells.ranges)


def test_pdf_export_is_a_pdf():
    data = render_pdf(sample_gen(), "faculty")
    assert data.startswith(b"%PDF") and len(data) > 1000


def test_empty_export_does_not_crash():
    gen = sample_gen()
    gen["sessions"] = []
    assert render_xlsx(gen)[:2] == b"PK"
    assert render_pdf(gen).startswith(b"%PDF")


def test_lab_column_layout_in_grid():
    gen = sample_gen()
    sess = [s for s in gen["sessions"] if s["batch_id"] == 1]
    header, rows, spans = build_grid(gen["config"], sess, "batch", lab_column=True)
    assert len(header) == 1 + 2 * 5 and header[2].endswith("Lab")
    assert rows[1][4].startswith("CS102L") and rows[1][3] == ""   # lab session only in the Lab column
    assert (4, 1, 4, 2) in spans
